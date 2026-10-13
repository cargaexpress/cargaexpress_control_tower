"""Pruebas unitarias para el Simulador de Archivos Tipo C con Inyección de Defectos (F5, Sec 5.2 y 5.3)."""

import csv
from datetime import datetime
from pathlib import Path
import re

import openpyxl
import pytest

from simulators.type_c.catalogs import CARRIERS_REGISTRY, get_carrier_config
from simulators.type_c.cli import main as cli_main
from simulators.type_c.corruption import (
    inject_duplicate_rows,
    inject_header_and_summary_rows,
    inject_mixed_date_formats,
    inject_nonexistent_guides,
    inject_out_of_sequence_events,
)
from simulators.type_c.generator import TypeCSimulator, compute_file_sha256


class TestTypeCSimulator:
    """Suite de pruebas de integración y unitarias para el simulador Tipo C."""

    def test_reproducibility_byte_to_byte_csv(self, tmp_path: Path):
        """Validar que dos ejecuciones con la misma semilla generen exactamente el mismo archivo (hash idéntico)."""
        sim1 = TypeCSimulator(carrier=1, seed=42, error_rate=0.1)
        sim2 = TypeCSimulator(carrier=1, seed=42, error_rate=0.1)

        file1 = sim1.generate_file(output_dir=tmp_path / "run1", num_rows=80, file_format="csv")
        file2 = sim2.generate_file(output_dir=tmp_path / "run2", num_rows=80, file_format="csv")

        hash1 = compute_file_sha256(file1)
        hash2 = compute_file_sha256(file2)

        assert hash1 == hash2, "Las ejecuciones con la misma semilla deben producir el mismo hash SHA-256 byte a byte"
        assert file1.read_text(encoding="utf-8") == file2.read_text(encoding="utf-8")

    def test_reproducibility_different_seeds(self, tmp_path: Path):
        """Semillas diferentes deben producir archivos diferentes."""
        sim_a = TypeCSimulator(carrier=1, seed=42, error_rate=0.1)
        sim_b = TypeCSimulator(carrier=1, seed=99, error_rate=0.1)

        file_a = sim_a.generate_file(output_dir=tmp_path / "run_a", num_rows=50, file_format="csv")
        file_b = sim_b.generate_file(output_dir=tmp_path / "run_b", num_rows=50, file_format="csv")

        assert compute_file_sha256(file_a) != compute_file_sha256(file_b)

    def test_clean_generation_without_errors(self, tmp_path: Path):
        """Con error_rate=0, el archivo se genera limpio, sin cabeceras repetidas ni filas de resumen."""
        sim = TypeCSimulator(carrier=1, seed=42, error_rate=0.0)
        csv_file = sim.generate_file(output_dir=tmp_path, num_rows=50, file_format="csv")

        with open(csv_file, mode="r", encoding="utf-8") as f:
            reader = list(csv.reader(f))

        header = reader[0]
        data_rows = reader[1:]

        assert header == sim.config.columns
        assert len(data_rows) == 50

        # No debe haber filas repetidas de cabecera
        for row in data_rows:
            assert row != header, "No debe haber encabezados repetidos en generación limpia"
            assert not any("TOTAL REGISTROS" in str(col) for col in row), "No debe haber fila de total"
            assert not any(str(col).startswith("FAKE_WAYBILL") for col in row), "No debe haber guías ficticias"

            # Validar formato estándar ISO de fecha: YYYY-MM-DD HH:MM:SS
            ts_str = row[sim.config.columns.index(sim.config.timestamp_col)]
            assert re.match(r"^\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2}$", ts_str), f"Fecha debe ser ISO estándar: {ts_str}"

    def test_corrupted_structure_injection(self, tmp_path: Path):
        """Con error_rate > 0, el archivo resultante contiene encabezados repetidos y pie de resumen."""
        sim = TypeCSimulator(carrier=1, seed=42, error_rate=0.2)
        csv_file = sim.generate_file(output_dir=tmp_path, num_rows=60, file_format="csv")

        with open(csv_file, mode="r", encoding="utf-8") as f:
            content = f.read()

        # Verificar presencia de fila de resumen al pie
        assert "TOTAL REGISTROS:" in content
        assert "--- FIN DE LOTE ---" in content

        # Verificar presencia de encabezados repetidos en el cuerpo
        lines = content.strip().split("\n")
        header_line = ",".join(sim.config.columns)
        assert lines[0] == header_line
        repeated_headers = [line for line in lines[1:] if line == header_line]
        assert len(repeated_headers) >= 1, "Debe contener al menos un encabezado repetido en el cuerpo"

    def test_mixed_date_formats_presence(self, tmp_path: Path):
        """Con error_rate > 0, debe detectarse la presencia de formatos mixtos DD/MM/YYYY y MM/DD/YYYY."""
        sim = TypeCSimulator(carrier=1, seed=42, error_rate=0.3)
        events = sim.generate_events(num_rows=100)

        date_col = sim.config.timestamp_col
        timestamps = [row[date_col] for row in events if isinstance(row.get(date_col), str)]

        # Patrón mixto: DD/MM/YYYY HH:MM o MM/DD/YYYY HH:MM
        slash_dates = [ts for ts in timestamps if "/" in ts]
        assert len(slash_dates) > 0, "Debe haber fechas con formato intercalado de barra inclinada"

        # Validar que los patrones cumplen con día/mes/año hora:minuto
        for dt_str in slash_dates:
            assert re.match(r"^\d{2}/\d{2}/\d{4} \d{2}:\d{2}$", dt_str), f"Fecha corrupta debe coincidir con formato: {dt_str}"

    def test_fake_guides_and_duplicates_injected(self):
        """Validar presencia de guías inexistentes y registros duplicados con error_rate > 0."""
        sim = TypeCSimulator(carrier=1, seed=42, error_rate=0.2)
        events = sim.generate_events(num_rows=80)

        # 1. Guías ficticias
        fake_guides = [row for row in events if str(row.get("waybill_no", "")).startswith("FAKE_WAYBILL_")]
        assert len(fake_guides) > 0, "Debe haber guías inexistentes inyectadas"

        # 2. Filas duplicadas
        # Excluir cabeceras repetidas o pie
        tracking_values = [row.get("waybill_no") for row in events if "TOTAL" not in str(row.get("waybill_no"))]
        assert len(tracking_values) > len(set(tracking_values)), "Debe haber guías repetidas/duplicadas en los eventos"

    def test_xlsx_format_generation(self, tmp_path: Path):
        """Validar generación de archivo en formato Excel (.xlsx)."""
        sim = TypeCSimulator(carrier=2, seed=42, error_rate=0.1)
        xlsx_file = sim.generate_file(output_dir=tmp_path, num_rows=40, file_format="xlsx")

        assert xlsx_file.exists()
        assert xlsx_file.suffix == ".xlsx"

        wb = openpyxl.load_workbook(xlsx_file)
        assert sim.config.code in wb.sheetnames
        ws = wb[sim.config.code]

        rows = list(ws.iter_rows(values_only=True))
        assert len(rows) > 40
        assert list(rows[0]) == sim.config.columns

    def test_support_for_all_three_carriers(self, tmp_path: Path):
        """Verificar que las 3 transportadoras se puedan simular correctamente con sus catálogos."""
        for c_id in [1, 2, 3]:
            config = get_carrier_config(c_id)
            assert len(config.status_catalog) >= 6, f"Carrier {c_id} debe tener al menos 6 estados en catálogo"

            sim = TypeCSimulator(carrier=c_id, seed=100 + c_id, error_rate=0.0)
            target_path = sim.generate_file(output_dir=tmp_path / f"c_{c_id}", num_rows=30)
            assert target_path.exists()

    def test_cli_execution(self, tmp_path: Path, capsys):
        """Verificar la ejecución del comando CLI con argumentos."""
        out_dir = tmp_path / "cli_output"
        exit_code = cli_main([
            "--seed", "42",
            "--transportadora", "1",
            "--output-dir", str(out_dir),
            "--rows", "25",
            "--error-rate", "0.1",
            "--format", "csv",
        ])

        assert exit_code == 0
        captured = capsys.readouterr()
        assert "SIMULADOR TIPO C - GENERACIÓN COMPLETADA CON ÉXITO" in captured.out
        assert "SHA-256 Hash" in captured.out
        assert any(out_dir.glob("*.csv"))
